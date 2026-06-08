from __future__ import annotations

import asyncio
import base64
import concurrent.futures
import io
import json
import logging
import struct
import zlib
from datetime import datetime, timezone
from typing import Any
from uuid import UUID

from app.core.classification import ClassificationLevel
from app.services.signing import report_signer

logger = logging.getLogger(__name__)

CLASSIFICATION_BANNERS: dict[int, str] = {
    ClassificationLevel.PUBLIC: "PUBLIC",
    ClassificationLevel.RESTRICTED: "RESTRICTED",
    ClassificationLevel.CONFIDENTIAL: "CONFIDENTIAL",
    ClassificationLevel.SECRET: "SECRET",
}

_WATERMARK_MAGIC = b"GIWM"
_WATERMARK_VERSION = 1


def _resolve_classification(classification: str | int) -> int:
    if isinstance(classification, int):
        return classification
    return int(ClassificationLevel.from_any(classification))


def _classification_label(level: int) -> str:
    return CLASSIFICATION_BANNERS.get(level, "UNKNOWN")


def _fetch_report_data(report_id: UUID) -> dict[str, Any]:
    from app.db.session import AsyncSessionLocal
    from app.models.reports import DailyReport, PremiumReport
    from sqlalchemy.future import select

    async def _load() -> dict[str, Any]:
        async with AsyncSessionLocal() as session:
            result = await session.execute(
                select(DailyReport).where(DailyReport.id == report_id)
            )
            report = result.scalars().first()
            if report:
                return {
                    "id": str(report.id),
                    "type": "daily",
                    "country_id": str(report.country_id),
                    "report_date": report.report_date.isoformat() if report.report_date else None,
                    "executive_summary": report.executive_summary or "",
                    "content_json": report.content_json or "{}",
                    "classification": int(report.classification or 0),
                    "tlp": report.tlp,
                    "org_id": str(report.org_id) if report.org_id else None,
                    "signature": report.signature,
                    "signed_at": report.signed_at.isoformat() if report.signed_at else None,
                    "signature_fingerprint": report.signature_fingerprint,
                }

            result = await session.execute(
                select(PremiumReport).where(PremiumReport.id == report_id)
            )
            report = result.scalars().first()
            if report:
                return {
                    "id": str(report.id),
                    "type": "premium",
                    "client_id": str(report.client_id),
                    "request_topic": report.request_topic or "",
                    "content_markdown": report.content_markdown or "",
                    "classification": int(report.classification or 0),
                    "tlp": report.tlp,
                    "org_id": str(report.org_id) if report.org_id else None,
                    "signature": report.signature,
                    "signed_at": report.signed_at.isoformat() if report.signed_at else None,
                    "signature_fingerprint": report.signature_fingerprint,
                }

            raise ValueError(f"Report {report_id} not found")

    try:
        asyncio.get_running_loop()
        with concurrent.futures.ThreadPoolExecutor() as pool:
            return pool.submit(asyncio.run, _load()).result()
    except RuntimeError:
        return asyncio.run(_load())


def _build_watermark_bytes(user_id: UUID, timestamp: str, classification: int) -> bytes:
    payload = json.dumps(
        {
            "user_id": str(user_id),
            "timestamp": timestamp,
            "classification": classification,
        },
        sort_keys=True,
        separators=(",", ":"),
    ).encode("utf-8")
    compressed = zlib.compress(payload)
    header = _WATERMARK_MAGIC + struct.pack("!BI", _WATERMARK_VERSION, len(compressed))
    return header + compressed


def _parse_content_json(raw: str | dict | None) -> Any:
    if isinstance(raw, dict):
        return raw
    if not raw:
        return {}
    try:
        return json.loads(raw)
    except (json.JSONDecodeError, TypeError):
        return {}


def _render_html_for_pdf(report: dict[str, Any], classification_label: str) -> str:
    title = (
        report.get("executive_summary", "")[:120]
        or report.get("request_topic", "Intelligence Report")
    )
    body_content = report.get("executive_summary", "") or report.get("content_markdown", "")
    content_structured = _parse_content_json(report.get("content_json"))

    sections_html = ""
    if isinstance(content_structured, dict):
        for section_key, section_value in content_structured.items():
            sections_html += f"<h2>{section_key}</h2>"
            if isinstance(section_value, str):
                sections_html += f"<p>{section_value}</p>"
            elif isinstance(section_value, list):
                sections_html += "<ul>"
                for item in section_value:
                    sections_html += f"<li>{item}</li>"
                sections_html += "</ul>"
            else:
                sections_html += f"<p>{json.dumps(section_value)}</p>"

    signature_block = ""
    if report.get("signature"):
        signature_block = (
            f'<div style="margin-top:30px;border-top:1px solid #ccc;padding-top:10px;font-size:9px;color:#666;">'
            f"<p>Signature Fingerprint: {report.get('signature_fingerprint', 'N/A')}</p>"
            f"<p>Signed At: {report.get('signed_at', 'N/A')}</p>"
            f"</div>"
        )

    return f"""<!DOCTYPE html>
<html>
<head><meta charset="utf-8"><title>{title}</title></head>
<body style="font-family:serif;margin:0;padding:0;">
<div style="background:#000;color:#fff;text-align:center;padding:8px;font-weight:bold;font-size:14px;letter-spacing:3px;">
{classification_label}
</div>
<div style="padding:20px;">
<h1 style="text-align:center;">Global Intelligence Platform</h1>
<h2 style="text-align:center;color:#555;">{title}</h2>
<p style="text-align:center;color:#888;font-size:11px;">
Report ID: {report.get('id', 'N/A')} | Date: {report.get('report_date', 'N/A')} | TLP: {report.get('tlp', 'N/A')}
</p>
<hr/>
<div style="margin-top:20px;">
<h3>Executive Summary</h3>
<p>{body_content}</p>
</div>
{sections_html}
{signature_block}
</div>
<div style="background:#000;color:#fff;text-align:center;padding:8px;font-weight:bold;font-size:14px;letter-spacing:3px;">
{classification_label}
</div>
</body>
</html>"""


def _sign_content(content: bytes) -> dict[str, str]:
    import hashlib

    content_hash = hashlib.sha256(content).hexdigest()

    if not report_signer.available:
        return {"signature": "", "fingerprint": "", "content_hash": content_hash}

    class _ExportPayload:
        def __init__(self, data_hash: str) -> None:
            self.id = "export"
            self.classification = ClassificationLevel.PUBLIC
            self.tlp = "TLP:CLEAR"
            self.org_id = None
            self.executive_summary = ""
            self.content_json = data_hash
            self.content_markdown = ""
            self.report_date = datetime.now(timezone.utc)
            self.created_at = datetime.now(timezone.utc)

    try:
        payload = _ExportPayload(content_hash)
        signature = report_signer.sign(payload)
        fingerprint = report_signer.fingerprint
        return {"signature": signature, "fingerprint": fingerprint, "content_hash": content_hash}
    except Exception as exc:
        logger.error("Failed to sign export: %s", exc)
        return {"signature": "", "fingerprint": "", "content_hash": content_hash}


async def export_pdf(report_id: UUID, classification: str) -> bytes:
    report = _fetch_report_data(report_id)
    cls_level = _resolve_classification(classification)
    label = _classification_label(cls_level)

    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        from reportlab.lib.units import inch
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_CENTER

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            topMargin=1.2 * inch,
            bottomMargin=1.2 * inch,
        )

        styles = getSampleStyleSheet()
        banner_style = ParagraphStyle(
            "ClassificationBanner",
            parent=styles["Heading1"],
            alignment=TA_CENTER,
            fontSize=16,
            textColor=colors.white,
            backColor=colors.black,
            spaceAfter=12,
            spaceBefore=0,
            leading=20,
        )
        body_style = styles["BodyText"]
        heading_style = styles["Heading2"]

        elements: list[Any] = []
        elements.append(Paragraph(label, banner_style))
        elements.append(Spacer(1, 0.3 * inch))
        elements.append(Paragraph("Global Intelligence Platform", styles["Title"]))
        elements.append(Spacer(1, 0.2 * inch))
        elements.append(
            Paragraph(
                f"Report ID: {report.get('id', 'N/A')} | "
                f"Date: {report.get('report_date', 'N/A')} | "
                f"TLP: {report.get('tlp', 'N/A')}",
                styles["Normal"],
            )
        )
        elements.append(Spacer(1, 0.3 * inch))

        exec_summary = report.get("executive_summary", "") or report.get("content_markdown", "")
        if exec_summary:
            elements.append(Paragraph("Executive Summary", heading_style))
            elements.append(Paragraph(exec_summary, body_style))
            elements.append(Spacer(1, 0.2 * inch))

        content_structured = _parse_content_json(report.get("content_json"))
        if isinstance(content_structured, dict):
            for section_key, section_value in content_structured.items():
                elements.append(Paragraph(str(section_key), heading_style))
                if isinstance(section_value, str):
                    elements.append(Paragraph(section_value, body_style))
                elif isinstance(section_value, list):
                    for item in section_value:
                        elements.append(Paragraph(f"\u2022 {item}", body_style))
                else:
                    elements.append(Paragraph(json.dumps(section_value), body_style))
                elements.append(Spacer(1, 0.15 * inch))

        if report.get("signature"):
            elements.append(Spacer(1, 0.3 * inch))
            elements.append(Paragraph("Digital Signature", heading_style))
            elements.append(
                Paragraph(f"Fingerprint: {report.get('signature_fingerprint', 'N/A')}", styles["Normal"])
            )
            elements.append(
                Paragraph(f"Signed At: {report.get('signed_at', 'N/A')}", styles["Normal"])
            )

        elements.append(Spacer(1, 0.3 * inch))
        elements.append(Paragraph(label, banner_style))

        doc.build(elements)
        pdf_bytes = buffer.getvalue()
        buffer.close()
        return pdf_bytes

    except ImportError:
        logger.warning("reportlab not installed, falling back to HTML-based PDF")
        return _render_html_for_pdf(report, label).encode("utf-8")


async def export_docx(report_id: UUID) -> bytes:
    report = _fetch_report_data(report_id)
    cls_level = int(report.get("classification", 0))
    label = _classification_label(cls_level)

    try:
        from docx import Document
        from docx.shared import Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        doc = Document()

        header = doc.sections[0].header
        header_para = header.paragraphs[0]
        header_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        header_run = header_para.add_run(label)
        header_run.bold = True
        header_run.font.size = Pt(14)

        footer = doc.sections[0].footer
        footer_para = footer.paragraphs[0]
        footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer_run = footer_para.add_run(label)
        footer_run.bold = True
        footer_run.font.size = Pt(14)

        title = doc.add_heading("Global Intelligence Platform", level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        subtitle_text = (
            f"Report ID: {report.get('id', 'N/A')} | "
            f"Date: {report.get('report_date', 'N/A')} | "
            f"TLP: {report.get('tlp', 'N/A')}"
        )
        subtitle = doc.add_paragraph(subtitle_text)
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

        doc.add_paragraph()

        exec_summary = report.get("executive_summary", "") or report.get("content_markdown", "")
        if exec_summary:
            doc.add_heading("Executive Summary", level=1)
            doc.add_paragraph(exec_summary)

        content_structured = _parse_content_json(report.get("content_json"))
        if isinstance(content_structured, dict):
            for section_key, section_value in content_structured.items():
                doc.add_heading(str(section_key), level=1)
                if isinstance(section_value, str):
                    doc.add_paragraph(section_value)
                elif isinstance(section_value, list):
                    for item in section_value:
                        doc.add_paragraph(str(item), style="List Bullet")
                else:
                    doc.add_paragraph(json.dumps(section_value))

        if report.get("signature"):
            doc.add_paragraph()
            doc.add_heading("Digital Signature", level=1)
            doc.add_paragraph(f"Fingerprint: {report.get('signature_fingerprint', 'N/A')}")
            doc.add_paragraph(f"Signed At: {report.get('signed_at', 'N/A')}")

        buffer = io.BytesIO()
        doc.save(buffer)
        docx_bytes = buffer.getvalue()
        buffer.close()
        return docx_bytes

    except ImportError:
        logger.warning("python-docx not installed, returning JSON fallback for DOCX export")
        fallback = json.dumps(
            {"error": "python-docx not installed", "report": report},
            default=str,
        )
        return fallback.encode("utf-8")


async def export_json(report_id: UUID) -> bytes:
    report = _fetch_report_data(report_id)
    parsed_content = _parse_content_json(report.get("content_json"))

    export_payload = {
        "metadata": {
            "report_id": report.get("id"),
            "report_type": report.get("type"),
            "classification": _classification_label(int(report.get("classification", 0))),
            "classification_level": int(report.get("classification", 0)),
            "tlp": report.get("tlp"),
            "org_id": report.get("org_id"),
            "report_date": report.get("report_date"),
            "exported_at": datetime.now(timezone.utc).isoformat(),
            "platform": "Global Intelligence Platform",
        },
        "content": {
            "executive_summary": report.get("executive_summary"),
            "structured_data": parsed_content,
            "markdown": report.get("content_markdown"),
        },
        "signature": {
            "value": report.get("signature"),
            "fingerprint": report.get("signature_fingerprint"),
            "signed_at": report.get("signed_at"),
            "algorithm": "Ed25519",
        },
    }

    return json.dumps(
        export_payload, indent=2, sort_keys=False, default=str, ensure_ascii=False
    ).encode("utf-8")


async def export_markdown(report_id: UUID) -> str:
    report = _fetch_report_data(report_id)
    cls_label = _classification_label(int(report.get("classification", 0)))

    lines: list[str] = []
    lines.append(f"# [{cls_label}]")
    lines.append("")
    lines.append("# Global Intelligence Platform")
    lines.append("")
    lines.append(f"**Report ID:** {report.get('id', 'N/A')}  ")
    lines.append(f"**Date:** {report.get('report_date', 'N/A')}  ")
    lines.append(f"**TLP:** {report.get('tlp', 'N/A')}  ")
    lines.append(f"**Classification:** {cls_label}  ")
    lines.append("")
    lines.append("---")
    lines.append("")

    exec_summary = report.get("executive_summary", "") or report.get("content_markdown", "")
    if exec_summary:
        lines.append("## Executive Summary")
        lines.append("")
        lines.append(exec_summary)
        lines.append("")

    content_structured = _parse_content_json(report.get("content_json"))
    if isinstance(content_structured, dict):
        for section_key, section_value in content_structured.items():
            lines.append(f"## {section_key}")
            lines.append("")
            if isinstance(section_value, str):
                lines.append(section_value)
            elif isinstance(section_value, list):
                for item in section_value:
                    lines.append(f"- {item}")
            else:
                lines.append(f"```json\n{json.dumps(section_value, indent=2)}\n```")
            lines.append("")

    if report.get("signature"):
        lines.append("---")
        lines.append("")
        lines.append("## Digital Signature")
        lines.append("")
        lines.append("- **Algorithm:** Ed25519")
        lines.append(f"- **Fingerprint:** {report.get('signature_fingerprint', 'N/A')}")
        lines.append(f"- **Signed At:** {report.get('signed_at', 'N/A')}")
        lines.append("")

    lines.append(f"# [{cls_label}]")
    lines.append("")

    return "\n".join(lines)


async def add_watermark(content: bytes, user_id: UUID, timestamp: str) -> bytes:
    classification = ClassificationLevel.PUBLIC
    watermark_data = _build_watermark_bytes(user_id, timestamp, int(classification))

    container = io.BytesIO()
    container.write(content)
    container.write(b"\x00")
    container.write(watermark_data)

    return container.getvalue()


async def sign_export(content: bytes) -> dict[str, str]:
    return _sign_content(content)


__all__ = [
    "export_pdf",
    "export_docx",
    "export_json",
    "export_markdown",
    "add_watermark",
    "sign_export",
]
